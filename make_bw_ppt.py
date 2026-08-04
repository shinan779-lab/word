from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE

# Transcribed from the supplied figure. O/Y/G correspond to original orange/yellow/green.
rows = [
"OOOOOOGGGYG", "GGGGGGOOYGG", "GGGGGGGGGOO",
"OOOOOOGGGYG", "GGGGGGOOYGG", "GYGGGGGGGOO",
"OOOOOOGGGYY", "GGGGGGOOOGY", "GYGGGGYOYOO",
"OOOOOOGGGYY", "GGGGGGOOYGO", "YYGGGGYGOOO",
"OOOOOOYYGYY", "YYYGGGOYOYY", "GYGGGGOOOOO",
"OOOOOOGGGYG", "GGGGGGOOYGG", "GYGGGGGGGOO",
]
groups = ["SPEI−30", "SPEI−60", "SPEI−90", "SPEI−180", "SPEI−270", "SPEI−360"]
cols = ["ET", "Prec", "PET", "Tmax", "Tmin", "Tmean", "AO", "NAO", "PNA", "SM", "Wind"]

prs = Presentation()
prs.slide_width = Inches(9)
prs.slide_height = Inches(11.8)
slide = prs.slides.add_slide(prs.slide_layouts[6])

BLACK=RGBColor(0,0,0); WHITE=RGBColor(255,255,255); LGRAY=RGBColor(205,205,205); MGRAY=RGBColor(150,150,150)

def rect(x,y,w,h,fill=WHITE,line=BLACK,lw=0.8):
    s=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb=fill
    s.line.color.rgb=line; s.line.width=Pt(lw)
    return s

def label(shape,text,size=13,bold=False,rotation=0,font="Arial"):
    tf=shape.text_frame; tf.clear(); tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=tf.paragraphs[0]; p.alignment=PP_ALIGN.CENTER
    r=p.add_run(); r.text=text; r.font.name=font; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=BLACK
    shape.rotation=rotation

# Layout
left=0.72; top=0.28; cw=0.62; rh=0.48; groupw=0.72; pcw=0.63
matrix_x=left+groupw
# White background and subtle outer frame
bg=rect(0.03,0.03,8.94,11.73,WHITE,BLACK,1.2)
slide.shapes._spTree.remove(bg._element); slide.shapes._spTree.insert(2,bg._element)

# Group labels and PC labels
for g,name in enumerate(groups):
    y=top+g*3*rh
    s=rect(left,y,groupw,3*rh,LGRAY,BLACK,0.9); label(s,name,15,False,270,"Times New Roman")
    for rr in range(3):
        p=rect(matrix_x+11*cw,y+rr*rh,pcw,rh,LGRAY,BLACK,0.8)
        label(p,f"PC{rr+1}",13,False,0,"Times New Roman")

# circles: three grayscale encodings, all individual editable shapes
for r,row in enumerate(rows):
    for c,kind in enumerate(row):
        d=0.35; x=matrix_x+c*cw+(cw-d)/2; y=top+r*rh+(rh-d)/2
        sh=slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
        sh.line.color.rgb=BLACK
        if kind=='O':
            sh.fill.solid(); sh.fill.fore_color.rgb=BLACK; sh.line.width=Pt(1.0)
        elif kind=='Y':
            sh.fill.solid(); sh.fill.fore_color.rgb=WHITE; sh.line.width=Pt(2.0)
        else:
            sh.fill.solid(); sh.fill.fore_color.rgb=LGRAY; sh.line.width=Pt(1.4); sh.line.dash_style=MSO_LINE_DASH_STYLE.DASH
        # Keep each marker as a native editable PowerPoint shape.
        # The three fill/outline treatments remain distinct in grayscale printing.

# bottom headers
hy=top+18*rh
s=rect(left,hy,groupw,rh,LGRAY,BLACK,0.8); label(s,"n",13,False,0,"Times New Roman")
for c,name in enumerate(cols):
    h=rect(matrix_x+c*cw,hy,cw,rh,LGRAY,BLACK,0.8); label(h,name,11.5,False,0,"Times New Roman")
h=rect(matrix_x+11*cw,hy,pcw,rh,LGRAY,BLACK,0.8); label(h,"PCs",12,False,0,"Times New Roman")

# Legend with explicit original-color mapping
legend_y=hy+0.68
items=[('O','原橙色：黑色实心'),('Y','原黄色：白色空心'),('G','原绿色：灰色＋虚线边框')]
for i,(kind,txt) in enumerate(items):
    x=0.85+i*2.65
    d=.34
    sh=slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(legend_y), Inches(d), Inches(d))
    sh.line.color.rgb=BLACK
    if kind=='O': sh.fill.solid(); sh.fill.fore_color.rgb=BLACK; sh.line.width=Pt(1)
    elif kind=='Y': sh.fill.solid(); sh.fill.fore_color.rgb=WHITE; sh.line.width=Pt(2)
    else: sh.fill.solid(); sh.fill.fore_color.rgb=LGRAY; sh.line.width=Pt(1.4); sh.line.dash_style=MSO_LINE_DASH_STYLE.DASH
    tb=slide.shapes.add_textbox(Inches(x+.43), Inches(legend_y-.04), Inches(2.05), Inches(.43))
    label(tb,txt,10.5,False,0,"Microsoft YaHei")

# Small editability note
note=slide.shapes.add_textbox(Inches(.7), Inches(11.25), Inches(7.7), Inches(.28))
label(note,"所有圆点、文字与表格均为独立可编辑的 PowerPoint 元素",9,False,0,"Microsoft YaHei")

prs.save('球_黑白可编辑版.pptx')
print('saved 球_黑白可编辑版.pptx', len(slide.shapes), 'editable shapes')
